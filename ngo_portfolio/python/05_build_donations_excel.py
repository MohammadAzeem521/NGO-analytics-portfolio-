"""
05_build_donations_excel.py
----------------------------------------------------------------------
Builds a SEPARATE professional Excel workbook (Donations_Analytics.xlsx)
summarizing the 6,000-row synthetic donations dataset — pivot-style
summary tables, KPI cards, and charts — to demonstrate Excel analysis on
a realistically sized dataset (distinct from the 10-row NGO research
workbook).
----------------------------------------------------------------------
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

NAVY = "1F3864"
BLUE = "2E5395"
LIGHT_BLUE = "D9E2F3"
ORANGE = "ED7D31"
GREEN = "70AD47"
WHITE = "FFFFFF"
GREY = "404040"
ALT_ROW = "F2F5FB"
FONT = "Calibri"

DATA = "../data"
fact = pd.read_csv(f"{DATA}/fact_donations.csv")
dim_ngo = pd.read_csv(f"{DATA}/dim_ngo.csv")
dim_date = pd.read_csv(f"{DATA}/dim_date.csv", parse_dates=["date"])

df = fact.merge(dim_ngo, on="ngo_id").merge(dim_date, on="date_id")

wb = Workbook()

def header(ws, cell, text, span, fill=NAVY, size=14):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=size, bold=True, color=WHITE)
    ws[cell].fill = PatternFill("solid", fgColor=fill)
    ws[cell].alignment = Alignment(horizontal="center", vertical="center")
    row = int(cell[1:]); col = cell[0]
    from openpyxl.utils import column_index_from_string
    ws.merge_cells(start_row=row, start_column=column_index_from_string(col),
                    end_row=row, end_column=column_index_from_string(col)+span-1)
    ws.row_dimensions[row].height = 26

def write_table(ws, start_row, start_col, headers, rows, col_widths=None):
    for j, h in enumerate(headers):
        c = ws.cell(row=start_row, column=start_col+j, value=h)
        c.font = Font(name=FONT, bold=True, color=WHITE, size=10)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if col_widths:
            ws.column_dimensions[get_column_letter(start_col+j)].width = col_widths[j]
    thin = Side(style="thin", color="C9C9C9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = ws.cell(row=start_row+1+i, column=start_col+j, value=val)
            c.border = border
            c.font = Font(name=FONT, size=10, color=GREY)
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ALT_ROW)
    return start_row + 1 + len(rows)

# ============================================================
# SHEET 1: SUMMARY / KPIs
# ============================================================
ws1 = wb.active
ws1.title = "Executive Summary"
ws1.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [26, 18, 18, 18, 18]):
    ws1.column_dimensions[col].width = w
header(ws1, "A1", "NGO DONATIONS ANALYTICS — EXECUTIVE SUMMARY", 5, NAVY, 16)
ws1["A2"] = "Simulated donations dataset (2021-2025) | For SQL / Python / Power BI / Excel portfolio project"
ws1["A2"].font = Font(name=FONT, italic=True, size=9.5, color=GREY)
ws1.merge_cells("A2:E2")

kpis = [
    ("Total Donations (INR)", f"₹{df['donation_amount_inr'].sum():,.0f}"),
    ("Total Donations Count", f"{len(df):,}"),
    ("Average Donation (INR)", f"₹{df['donation_amount_inr'].mean():,.0f}"),
    ("Recurring Donor Rate", f"{df['is_recurring'].mean()*100:.1f}%"),
    ("Active NGOs", f"{df['ngo_id'].nunique()}"),
]
r = 4
for label, val in kpis:
    ws1.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True, color=NAVY, size=11)
    ws1.cell(row=r, column=1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    vcell = ws1.cell(row=r, column=2, value=val)
    vcell.font = Font(name=FONT, bold=True, size=13, color=BLUE)
    r += 1

# ============================================================
# SHEET 2: DONATIONS BY NGO (pivot-style)
# ============================================================
ws2 = wb.create_sheet("Donations by NGO")
ws2.sheet_view.showGridLines = False
header(ws2, "A1", "DONATIONS SUMMARY BY NGO", 5)
by_ngo = df.groupby("ngo_name").agg(
    total_donations=("donation_amount_inr", "sum"),
    num_donations=("donation_id", "count"),
    avg_donation=("donation_amount_inr", "mean"),
).reset_index().sort_values("total_donations", ascending=False)
rows = [[r["ngo_name"], round(r["total_donations"],2), int(r["num_donations"]), round(r["avg_donation"],2)]
        for _, r in by_ngo.iterrows()]
end_row = write_table(ws2, 3, 1, ["NGO Name", "Total Donations (INR)", "Number of Donations", "Avg Donation (INR)"],
                       rows, col_widths=[32, 20, 20, 18])
ws2.freeze_panes = "A4"

bar = BarChart()
bar.type = "bar"
bar.title = "Total Donations by NGO"
cats = Reference(ws2, min_col=1, min_row=4, max_row=end_row)
data = Reference(ws2, min_col=2, min_row=3, max_row=end_row)
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.height = 10
bar.width = 20
ws2.add_chart(bar, "F3")

# ============================================================
# SHEET 3: DONOR TYPE ANALYSIS
# ============================================================
ws3 = wb.create_sheet("Donor Type Analysis")
ws3.sheet_view.showGridLines = False
header(ws3, "A1", "DONOR TYPE ANALYSIS", 4)
by_donor = df.groupby("donor_type").agg(
    total=("donation_amount_inr", "sum"), count=("donation_id", "count")
).reset_index().sort_values("total", ascending=False)
rows = [[r["donor_type"], round(r["total"],2), int(r["count"])] for _, r in by_donor.iterrows()]
end_row3 = write_table(ws3, 3, 1, ["Donor Type", "Total Donations (INR)", "Number of Donations"], rows,
                        col_widths=[22, 20, 20])

pie = PieChart()
pie.title = "Donation Value Share by Donor Type"
cats3 = Reference(ws3, min_col=1, min_row=4, max_row=end_row3)
data3 = Reference(ws3, min_col=2, min_row=3, max_row=end_row3)
pie.add_data(data3, titles_from_data=True)
pie.set_categories(cats3)
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
pie.height = 9; pie.width = 14
ws3.add_chart(pie, "E3")

# ============================================================
# SHEET 4: MONTHLY TREND
# ============================================================
ws4 = wb.create_sheet("Monthly Trend")
ws4.sheet_view.showGridLines = False
header(ws4, "A1", "MONTHLY DONATION TREND", 3)
monthly = df.groupby(df["date"].dt.to_period("M"))["donation_amount_inr"].sum().reset_index()
monthly["date"] = monthly["date"].astype(str)
rows = [[m, round(v,2)] for m, v in zip(monthly["date"], monthly["donation_amount_inr"])]
end_row4 = write_table(ws4, 3, 1, ["Month", "Total Donations (INR)"], rows, col_widths=[16, 20])

line = LineChart()
line.title = "Monthly Donation Trend (2021-2025)"
cats4 = Reference(ws4, min_col=1, min_row=4, max_row=end_row4)
data4 = Reference(ws4, min_col=2, min_row=3, max_row=end_row4)
line.add_data(data4, titles_from_data=True)
line.set_categories(cats4)
line.height = 10; line.width = 24
ws4.add_chart(line, "D3")

wb.save("../excel/Donations_Analytics.xlsx")
print("Donations_Analytics.xlsx built successfully.")
